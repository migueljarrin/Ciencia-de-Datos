"""
Genera dbt/seeds/poblacion_cantonal.csv a partir de las Estimaciones y
Proyecciones de Población del INEC (Censo 2022, revisión 2024), cantonal.

Fuente: https://www.ecuadorencifras.gob.ec/proyecciones-poblacionales/
        .../censo_2022/revision_2024_areas/Cantonal.zip
Cada Excel provincial tiene una hoja por cantón y sexo (<canton>_h, <canton>_m)
con la fila "Total" por año (valores al 30 de junio). Se suma hombres + mujeres.

Uso:  python scripts/build_population_seed.py [carpeta_con_xlsx]
"""
import csv
import io
import re
import sys
import unicodedata
import urllib.request
import zipfile
from pathlib import Path

import openpyxl

URL = ("https://www.ecuadorencifras.gob.ec/documentos/web-inec/Poblacion_y_Demografia/"
       "Proyecciones_Poblacionales/censo_2022/revision_2024_areas/Cantonal.zip")
YEARS = range(2022, 2027)
OUT = Path(__file__).resolve().parents[1] / "dbt" / "seeds" / "poblacion_cantonal.csv"


def norm(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", s.lower())


def same_canton(label, official):
    """True si alguna palabra (>= 3 letras) de la etiqueta aparece en el nombre oficial."""
    words = [norm(w) for w in re.split(r"[_\s]+", str(label))]
    return any(len(w) >= 3 and w in norm(official) for w in words)


def workbooks(folder):
    if folder:
        for f in sorted(Path(folder).glob("*.xlsx")):
            yield f.name, openpyxl.load_workbook(f, read_only=True, data_only=True)
        return
    req = urllib.request.Request(URL, headers={"User-Agent": "Mozilla/5.0"})
    data = urllib.request.urlopen(req, timeout=120).read()
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        for name in sorted(n for n in z.namelist() if n.endswith(".xlsx")):
            yield name, openpyxl.load_workbook(io.BytesIO(z.read(name)), read_only=True, data_only=True)


def canton_index(ws):
    """Código y nombre de cada cantón desde la hoja Índice (bloque Hombres)."""
    codes = {}
    for row in ws.iter_rows(values_only=True):
        cells = [c for c in row if c is not None]
        if cells and str(cells[0]).startswith("Estimaciones y Proyecciones Mujeres"):
            break
        for i in range(len(cells) - 1):
            # El código viene como texto ('0101') o, a veces, como número (1109)
            raw = str(cells[i]).strip()
            if re.fullmatch(r"\d{3,4}", raw) and isinstance(cells[i + 1], str):
                codes[norm(cells[i + 1])] = (raw.zfill(4), cells[i + 1].strip())
    return codes


def sheet_totals(ws):
    title, years, totals = None, None, None
    for row in ws.iter_rows(values_only=True):
        cells = [c for c in row if c is not None]
        if not cells:
            continue
        if title is None and isinstance(cells[0], str) and cells[0].endswith("- Cantonal"):
            title = cells[0].replace("- Cantonal", "").strip()
        elif years is None and isinstance(cells[0], int) and cells[0] == 2010:
            years = cells
        elif years is not None and cells[0] == "Total":
            totals = dict(zip(years, cells[1:]))
            break
    return title, totals


def main():
    folder = sys.argv[1] if len(sys.argv) > 1 else None
    pop, warnings = {}, []
    for fname, wb in workbooks(folder):
        index = canton_index(wb[wb.sheetnames[0]])
        cantones = sorted(index.values())  # [(codigo, nombre)] en orden de código
        for suffix in ("_h", "_m"):
            sheets = [s for s in wb.sheetnames[1:] if s.endswith(suffix)]
            if len(sheets) != len(cantones):
                raise SystemExit(f"{fname}: {len(sheets)} hojas{suffix} vs {len(cantones)} cantones")
            # Las hojas siguen el orden de los códigos del índice. El título interno de la
            # hoja NO es confiable (p. ej. pablo_sexto_m dice "Portovelo"), así que se
            # empareja por posición y se verifica con el nombre de la hoja.
            for sheet, (code, name) in zip(sheets, cantones):
                if not same_canton(sheet[: -len(suffix)], name):
                    raise SystemExit(f"{fname}/{sheet}: no corresponde a {code} {name}")
                title, totals = sheet_totals(wb[sheet])
                if totals is None:
                    raise SystemExit(f"{fname}/{sheet}: no se encontró la fila Total")
                if title and not same_canton(title, name):
                    warnings.append(f"{fname}/{sheet}: título '{title}' no coincide con {name}")
                for y in YEARS:
                    pop[(code, name, y)] = pop.get((code, name, y), 0) + int(totals[y])
    OUT.parent.mkdir(parents=True, exist_ok=True)
    with open(OUT, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["canton_codigo", "canton_nombre_inec", "anio", "poblacion"])
        for (code, name, y), p in sorted(pop.items()):
            w.writerow([code, name, y, p])
    cantones = {c for c, _, _ in pop}
    for w in warnings:
        print("AVISO (error en la fuente INEC):", w)
    print(f"{OUT.name}: {len(cantones)} cantones x {len(YEARS)} años = {len(pop)} filas")


if __name__ == "__main__":
    main()
